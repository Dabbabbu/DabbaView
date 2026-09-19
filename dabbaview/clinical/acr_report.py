# Copyright (c) 2026 Park Seongho (Dabbabbu)
# This file is part of DabbaView, licensed under GPL-3.0.
# See LICENSE for details.
"""ACR 팬텀 QC 결과: 기록지 (PDF · Word · Excel), 기록 보관 (JSON), 날짜별 추세"""
import datetime
import json
import os

from . import acr

JUDGE = {True: "적합", False: "부적합", None: "-"}


# ═══ 기록 (추세 비교용) ═══

def history_path():
    override = os.environ.get("DABBAVIEW_ACR_HISTORY")
    if override:
        return override
    from ..library import library_path
    return os.path.join(os.path.dirname(library_path()), "acr_history.json")


def load_history(path=None):
    try:
        with open(path or history_path(), encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, list) else []
    except (OSError, ValueError):
        return []


def make_record(info, values, rows):
    """values: acr.compute() 결과 → JSON 저장 가능한 기록"""
    return {"date": info.get("date", ""), "hospital": info.get("hospital", ""),
            "unit": info.get("unit", ""), "scanner": info.get("scanner", ""),
            "field": info.get("field", ""), "coil": info.get("coil", ""),
            "values": {f"{t}|{s}": v for (t, s), v in values.items()},
            "rows": [list(r) for r in rows], "overall": acr.overall(rows),
            "saved": datetime.datetime.now().isoformat(timespec="seconds")}


def save_record(record, path=None):
    """같은 날짜·호기·장비 기록은 바꿔 씀. → 전체 기록"""
    path = path or history_path()
    items = [r for r in load_history(path)
             if (r.get("date"), r.get("unit"), r.get("scanner"))
             != (record["date"], record["unit"], record["scanner"])]
    items.append(record)
    items.sort(key=lambda r: (r.get("date", ""), r.get("unit", "")))
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=1)
    os.replace(tmp, path)
    return items


def delete_record(date, unit, scanner, path=None):
    path = path or history_path()
    items = [r for r in load_history(path) if (r.get("date"), r.get("unit"), r.get("scanner"))
             != (date, unit, scanner)]
    with open(path, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=1)
    return items


def _v(rec, key, field):
    v = rec.get("values", {}).get(key)
    return None if v is None else v.get(field)


def _geo_s5(rec):
    v = rec.get("values", {}).get("geometry|T1") or {}
    vals = [x for k, x in v.items() if k.startswith("S5")]
    return sum(vals) / len(vals) if vals else None


def _pos(rec, seq):
    v = rec.get("values", {}).get(f"position|{seq}") or {}
    return max((abs(x) for x in v.values()), default=None)


# (키, 이름, 값 꺼내기, 기준선 [(값, 설명)] 만들기)
TREND_METRICS = [
    ("geo_loc", "Localizer 길이 (mm)", lambda r: _v(r, "geometry|T1", "LOC"),
     lambda c: [(c["loc_nominal"] - c["geo_tol"], "하한"), (c["loc_nominal"] + c["geo_tol"], "상한")]),
    ("geo_s5", "Slice 5 지름 평균 (mm)", _geo_s5,
     lambda c: [(c["axial_nominal"] - c["geo_tol"], "하한"), (c["axial_nominal"] + c["geo_tol"], "상한")]),
    ("thk", "절편 두께 (mm)", None,
     lambda c: [(c["thk_nominal"] - c["thk_tol"], "하한"), (c["thk_nominal"] + c["thk_tol"], "상한")]),
    ("pos", "절편 위치 |차이| 최대 (mm)", None, lambda c: [(c["pos_max"], "상한")]),
    ("piu", "PIU (%)", None, lambda c: [(c["piu_min"], "하한")]),
    ("ghost", "고스팅 (%)", None, lambda c: [(c["ghost_max"], "상한")]),
    ("lc", "저대조도 스포크 합계", None, lambda c: [(c["lc_min"], "하한")]),
    ("res", "분해능 UL/LR 중 큰 값 (mm)", None, lambda c: [(c["res_max"], "상한")]),
]
SEQ_METRICS = {
    "thk": lambda r, s: _v(r, f"thickness|{s}", "thickness"),
    "pos": _pos,
    "piu": lambda r, s: _v(r, f"uniformity|{s}", "piu"),
    "ghost": lambda r, s: _v(r, f"ghosting|{s}", "ratio"),
    "lc": lambda r, s: _v(r, f"low_contrast|{s}", "total"),
    "res": lambda r, s: (lambda v: None if not v or v.get("ul") is None or v.get("lr") is None
                         else max(v["ul"], v["lr"]))(r.get("values", {}).get(f"resolution|{s}")),
}


def trend_series(history, key):
    """→ {이름: [(날짜, 값)]}"""
    metric = next(m for m in TREND_METRICS if m[0] == key)
    out = {}
    for rec in history:
        label = rec.get("unit") or rec.get("scanner") or ""
        if metric[2] is not None:
            series = {"T1": metric[2](rec)}
        else:
            series = {s: SEQ_METRICS[key](rec, s) for s in ("T1", "T2")}
        for seq, value in series.items():
            if value is not None:
                name = f"{label} {seq}".strip()
                out.setdefault(name, []).append((rec.get("date") or rec.get("saved", "")[:10], value))
    return out


def plot_trend(figure, history, key, criteria):
    metric = next(m for m in TREND_METRICS if m[0] == key)
    ax = figure.add_subplot(111)
    data = trend_series(history, key)
    for name, pts in sorted(data.items()):
        pts.sort()
        dates = [datetime.date.fromisoformat(d) if len(d) == 10 else d for d, _ in pts]
        ax.plot(dates, [v for _, v in pts], "o-", label=name)
    for value, text in metric[3](criteria):
        ax.axhline(value, color="#e5484d", linestyle="--", linewidth=1)
        ax.text(ax.get_xlim()[0], value, f" {text} {value:g}", color="#e5484d", fontsize=7, va="bottom")
    ax.set_title(metric[1])
    if data:
        ax.legend(fontsize=7)
    figure.autofmt_xdate()
    return ax


# ═══ 기록지 ═══

def half_of(date):
    try:
        d = datetime.date.fromisoformat(date)
        return f"{d.year}년 {'상' if d.month <= 6 else '하'}반기"
    except (TypeError, ValueError):
        return ""


def sheet_rows(rows):
    """evaluate() 행 → 기록지 표 [(항목, T1, T2, 기준, 판정)]"""
    by = {}
    for test, seq, measured, criterion, ok in rows:
        by.setdefault(test, {})[seq] = (measured, criterion, ok)
    out = []
    for test, name in acr.TESTS:
        items = by.get(test, {})
        if not items:
            out.append((name, "-", "-", "", "-"))
            continue
        t1 = items.get("T1", ("-", "", None))
        t2 = items.get("T2", ("-", "", None))
        judged = [x[2] for x in (t1, t2) if x[2] is not None]
        ok = None if not judged else all(judged)
        out.append((name, t1[0], t2[0], t1[1] or t2[1], JUDGE[ok]))
    return out


def info_rows(info):
    return [("병원", info.get("hospital", "")), ("연도 / 반기", half_of(info.get("date", ""))),
            ("장비명", info.get("scanner", "")), ("자장 세기", info.get("field", "")),
            ("호기", info.get("unit", "")), ("검사일", info.get("date", "")),
            ("코일", info.get("coil", "")), ("검사자", info.get("tester", ""))]


def notes_for(approximate):
    notes = ["DabbaView ACR Phantom QC 자동 분석 (ROI 위치는 사용자가 확인·수정할 수 있음).",
             "저대조도 스포크 수와 고대조도 분해능은 자동 추정 후 사용자가 확인한 값."]
    if approximate:
        notes.append("원본이 8비트 화면 캡처(JPEG)라 FOV로 픽셀 크기를 가정했고, 신호값은 표시 창(W/L)이 "
                     "적용된 값입니다 (L = W/2이면 비율 검사 PIU·고스팅은 유효).")
    return notes


TITLE = "MRI 팬텀 영상 정도관리 결과 (ACR 대형 팬텀)"


def write_report(fmt, path, info, rows, approximate=False, snapshots=(), steps=()):
    """fmt: 'pdf' | 'docx' | 'xlsx'. snapshots: [(제목, PNG 경로)] (PDF·Word에 붙임)
    steps: 측정 과정 [{"title", "images": [(설명, PNG)], "lines": [계산 단계]}] (Excel은 글만)"""
    writer = {"pdf": _write_pdf, "docx": _write_docx, "xlsx": _write_xlsx}[fmt]
    writer(path, info, rows, approximate, list(snapshots), list(steps))
    return path


STEPS_TITLE = "측정 과정 (콘솔 수동 절차와 같은 순서: ROI → W/L 조절 → 측정 → 계산 → 판정)"


def _image_size(p):
    from PIL import Image as PILImage
    with PILImage.open(p) as im:
        return im.size


def _overall_text(rows):
    ok = acr.overall(rows)
    return {True: "적합 (모든 항목 기준 만족)", False: "부적합 (기준 미달 항목 있음)", None: "-"}[ok]


def _write_pdf(path, info, rows, approximate, snapshots, steps=()):
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    from ..library_export import _pdf_font
    font = _pdf_font()
    base = ParagraphStyle("base", fontName=font, fontSize=9, leading=12)
    h1 = ParagraphStyle("h1", parent=base, fontSize=15, leading=20, alignment=1, spaceAfter=6)
    small = ParagraphStyle("small", parent=base, fontSize=7.5, leading=10, textColor=colors.grey)

    def para(text, style=base):
        return Paragraph(escape(str(text)), style)

    story = [para(TITLE, h1), Spacer(1, 2 * mm)]
    ir = info_rows(info)
    grid = [[para(ir[i][0]), para(ir[i][1]), para(ir[i + 1][0]), para(ir[i + 1][1])]
            for i in range(0, len(ir), 2)]
    t = Table(grid, colWidths=[25 * mm, 60 * mm, 25 * mm, 60 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                           ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF2F7")),
                           ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EEF2F7"))]))
    story += [t, Spacer(1, 4 * mm)]
    data = [[para(h) for h in ("검사 항목", "T1", "T2", "기준", "판정")]]
    styles = [("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
              ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE6F0")),
              ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for i, row in enumerate(sheet_rows(rows), start=1):
        data.append([para(c) for c in row])
        color = {"적합": "#DFF5E1", "부적합": "#FBE0E0"}.get(row[4])
        if color:
            styles.append(("BACKGROUND", (4, i), (4, i), colors.HexColor(color)))
    data.append([para("종합 판정"), para(_overall_text(rows)), "", "", ""])
    styles += [("SPAN", (1, len(data) - 1), (4, len(data) - 1)),
               ("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#F4F4F4"))]
    t = Table(data, colWidths=[38 * mm, 44 * mm, 44 * mm, 30 * mm, 14 * mm], repeatRows=1)
    t.setStyle(TableStyle(styles))
    story += [t, Spacer(1, 3 * mm)]
    story += [para("• " + n, small) for n in notes_for(approximate)]
    if snapshots:
        cells = [[Image(p, width=41 * mm, height=41 * mm), para(title, small)] for title, p in snapshots]
        rows_ = []
        for i in range(0, len(cells), 4):   # 한 줄에 4장 → A4 한 쪽
            chunk = cells[i:i + 4]
            rows_.append([c[0] for c in chunk] + [""] * (4 - len(chunk)))
            rows_.append([c[1] for c in chunk] + [""] * (4 - len(chunk)))
        story += [Spacer(1, 4 * mm), Table(rows_, colWidths=[44 * mm] * 4)]
    if steps:
        h2 = ParagraphStyle("h2", parent=base, fontSize=12, leading=16, spaceBefore=2, spaceAfter=3)
        h3 = ParagraphStyle("h3", parent=base, fontSize=10.5, leading=14, textColor=colors.HexColor("#1f3a5f"))
        story += [PageBreak(), para(STEPS_TITLE, h2)]
        full = 180 * mm
        for step in steps:
            block = [Spacer(1, 3 * mm), para(step["title"], h3), Spacer(1, 1.5 * mm)]
            imgs = step.get("images") or []
            if imgs:
                col = (full - 4 * mm * (len(imgs) - 1)) / len(imgs)
                cells, caps = [], []
                for caption, p in imgs:
                    w, h = _image_size(p)
                    iw = min(col, 90 * mm)
                    cells.append(Image(p, width=iw, height=iw * h / w))
                    caps.append(para(caption, small))
                t = Table([cells, caps], colWidths=[col + 4 * mm] * len(imgs))
                t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
                block.append(t)
            block += [para(line) for line in step.get("lines", [])]
            story.append(KeepTogether(block))
    SimpleDocTemplate(path, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm,
                      bottomMargin=14 * mm, title=TITLE).build(story)


def _write_docx(path, info, rows, approximate, snapshots, steps=()):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    def shade(cell, hex_color):
        tc = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tc.append(shd)

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "AppleGothic"
    style.font.size = Pt(9.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(15)
    title = doc.add_heading(TITLE, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ir = info_rows(info)
    t = doc.add_table(rows=len(ir) // 2, cols=4)
    t.style = "Table Grid"
    for i in range(0, len(ir), 2):
        r = t.rows[i // 2].cells
        r[0].text, r[1].text, r[2].text, r[3].text = ir[i][0], ir[i][1], ir[i + 1][0], ir[i + 1][1]
        shade(r[0], "EEF2F7")
        shade(r[2], "EEF2F7")
    doc.add_paragraph()
    sheet = sheet_rows(rows)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ("검사 항목", "T1", "T2", "기준", "판정")):
        c.text = h
        shade(c, "DDE6F0")
    for row in sheet:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = str(v)
        color = {"적합": "DFF5E1", "부적합": "FBE0E0"}.get(row[4])
        if color:
            shade(cells[4], color)
    last = t.add_row().cells
    last[0].text = "종합 판정"
    merged = last[1].merge(last[4])
    merged.text = _overall_text(rows)
    for n in notes_for(approximate):
        p = doc.add_paragraph("• " + n)
        p.runs[0].font.size = Pt(8)
    for title_text, p in snapshots:
        doc.add_paragraph(title_text).runs[0].font.size = Pt(8)
        doc.add_picture(p, width=Mm(70))
    if steps:
        doc.add_page_break()
        doc.add_heading(STEPS_TITLE, level=2)
        for step in steps:
            doc.add_heading(step["title"], level=3)
            imgs = step.get("images") or []
            if imgs:
                t = doc.add_table(rows=2, cols=len(imgs))
                width = Mm(min(85, 175 / len(imgs)))
                for i, (caption, p) in enumerate(imgs):
                    t.rows[0].cells[i].paragraphs[0].add_run().add_picture(p, width=width)
                    t.rows[1].cells[i].text = caption
            for line in step.get("lines", []):
                doc.add_paragraph(line).runs[0].font.size = Pt(9)
    doc.save(path)


def _write_xlsx(path, info, rows, approximate, snapshots, steps=()):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "ACR QC"
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = TITLE
    ws["A1"].font = Font(size=14, bold=True)
    ws.merge_cells("A1:E1")
    r = 3
    for key, value in info_rows(info):
        ws.cell(r, 1, key).fill = PatternFill("solid", fgColor="EEF2F7")
        ws.cell(r, 2, value)
        ws.cell(r, 1).border = ws.cell(r, 2).border = box
        r += 1
    r += 1
    for c, h in enumerate(("검사 항목", "T1", "T2", "기준", "판정"), start=1):
        cell = ws.cell(r, c, h)
        cell.fill = PatternFill("solid", fgColor="DDE6F0")
        cell.font = Font(bold=True)
        cell.border = box
    for row in sheet_rows(rows):
        r += 1
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            cell.border = box
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        color = {"적합": "DFF5E1", "부적합": "FBE0E0"}.get(row[4])
        if color:
            ws.cell(r, 5).fill = PatternFill("solid", fgColor=color)
    r += 1
    ws.cell(r, 1, "종합 판정").font = Font(bold=True)
    ws.cell(r, 2, _overall_text(rows))
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    r += 2
    for n in notes_for(approximate):
        ws.cell(r, 1, "• " + n).font = Font(size=8, color="777777")
        r += 1
    for col, width in zip("ABCDE", (26, 40, 40, 26, 10)):
        ws.column_dimensions[col].width = width
    if steps:
        ws2 = wb.create_sheet("측정 과정")
        ws2.append([STEPS_TITLE])
        ws2["A1"].font = Font(bold=True)
        for step in steps:
            ws2.append([])
            ws2.append([step["title"]])
            ws2.cell(ws2.max_row, 1).font = Font(bold=True)
            for line in step.get("lines", []):
                ws2.append([line])
        ws2.column_dimensions["A"].width = 140
    raw = wb.create_sheet("측정값")
    raw.append(["검사", "시퀀스", "측정값", "기준", "판정"])
    for test, seq, measured, criterion, ok in rows:
        raw.append([acr.TEST_NAMES.get(test, test), seq, measured, criterion, JUDGE[ok]])
    wb.save(path)
