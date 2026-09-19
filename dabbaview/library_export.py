# Copyright (c) 2026 Park Seongho (Dabbabbu)
# This file is part of DabbaView, licensed under GPL-3.0.
# See LICENSE for details.
"""Library 내보내기: PDF · Word · Excel · CSV · JSON · Markdown · PNG/JPEG (연락 시트/스터디별)

기록(record) 하나 = 스터디 하나: 환자·스터디 정보, 시리즈 목록, 메모, 태그, 썸네일, 측정/ROI 결과
HWP는 macOS에서 쓸 수 있는 라이브러리가 없어 (pyhwpx는 Windows 한글 자동화 전용) DOCX/PDF 안내
"""
import csv
import datetime
import json
import math
import os
import re
import sys

import numpy as np

from .library import _date, _html_to_markdown

INCLUDE_KEYS = [("patient", "환자 정보"), ("study", "스터디 정보 (날짜, Description)"),
                ("series", "시리즈 목록"), ("notes", "메모/코멘트"), ("tags", "태그"),
                ("thumbnail", "썸네일 이미지"), ("measurements", "측정/ROI 결과")]
FORMATS = [("pdf", "PDF (.pdf)"), ("docx", "Word (.docx)"), ("xlsx", "Excel (.xlsx)"),
           ("hwp", "HWP (한글)"), ("png_sheet", "PNG — 연락 시트 (한 장에 격자)"),
           ("png_each", "PNG — 스터디별 이미지"), ("jpg_sheet", "JPEG — 연락 시트"),
           ("jpg_each", "JPEG — 스터디별 이미지"), ("csv", "CSV (.csv)"),
           ("json", "JSON (전체 백업/복원)"), ("md", "Markdown (.md)")]
EXTENSIONS = {"pdf": "pdf", "docx": "docx", "xlsx": "xlsx", "csv": "csv", "json": "json",
              "md": "md", "png_sheet": "png", "jpg_sheet": "jpg", "hwp": "hwp"}
HWP_MESSAGE = ("HWP로 직접 저장하는 기능은 macOS에서 쓸 수 있는 라이브러리가 없어 지원하지 않습니다 "
               "(pyhwpx는 Windows의 한글 프로그램을 자동화하는 방식). 대신 Word(.docx)로 내보내면 "
               "한글에서 그대로 열어 HWP로 저장할 수 있습니다. PDF도 가능합니다.")


class UnsupportedFormat(Exception):
    pass


# ─── 글꼴 (한글) ───

def korean_font_path():
    for path in ("/System/Library/Fonts/Supplemental/AppleGothic.ttf",
                 "/Library/Fonts/NanumGothic.ttf", os.path.expanduser("~/Library/Fonts/NanumGothic.ttf"),
                 "C:/Windows/Fonts/malgun.ttf", "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"):
        if os.path.exists(path):
            return path
    return None


def _pil_font(size):
    from PIL import ImageFont
    path = korean_font_path()
    try:
        return ImageFont.truetype(path, size) if path else ImageFont.load_default()
    except OSError:
        return ImageFont.load_default()


# ─── 썸네일 ───

def render_thumbnail(arr, window=None, size=256):
    """2D 배열(또는 RGB) → 정사각형 PIL 이미지 (가운데 맞춤, 검은 배경)"""
    from PIL import Image
    a = np.asarray(arr, dtype=np.float32)
    if a.ndim == 3 and a.shape[-1] == 3:
        img = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8), "RGB")
    else:
        if window is None or not window[1]:
            lo, hi = np.percentile(a, [1, 99])
        else:
            lo, hi = window[0] - window[1] / 2, window[0] + window[1] / 2
        g = np.clip((a - lo) / max(hi - lo, 1e-6) * 255, 0, 255).astype(np.uint8)
        img = Image.fromarray(g, "L").convert("RGB")
    img.thumbnail((size, size))
    canvas = Image.new("RGB", (size, size), (0, 0, 0))
    canvas.paste(img, ((size - img.width) // 2, (size - img.height) // 2))
    return canvas


def thumbnails_dir(store):
    folder = os.path.join(os.path.dirname(store.path), "thumbnails")
    os.makedirs(folder, exist_ok=True)
    return folder


def save_series_thumbnail(store, study_uid, series):
    """불러온 시리즈의 중간 슬라이스 → 라이브러리 썸네일 (Library에 추가할 때)"""
    arr = series.get_pixel_array(series.num_slices // 2)
    if arr is None:
        return None
    try:
        window = series.get_default_window()
    except Exception:  # noqa: BLE001
        window = None
    path = os.path.join(thumbnails_dir(store), _safe(study_uid) + ".png")
    render_thumbnail(arr, window).save(path)
    return path


def find_thumbnail(store, uid, entry, limit=300):
    """저장된 썸네일, 없으면 스터디 폴더에서 이 스터디 파일 하나를 찾아 만듦 (없으면 None)"""
    path = os.path.join(thumbnails_dir(store), _safe(uid) + ".png")
    if os.path.exists(path):
        return path
    import pydicom
    from .dicom_loader import is_cloud_placeholder
    checked = 0
    for folder in entry.get("paths") or [entry.get("folder", "")]:
        if not folder or not os.path.isdir(folder):
            continue
        for root, _dirs, files in os.walk(folder):
            for name in sorted(files):
                f = os.path.join(root, name)
                if name.startswith(".") or name.endswith(".bak") or is_cloud_placeholder(f):
                    continue
                checked += 1
                if checked > limit:
                    return None
                try:
                    ds = pydicom.dcmread(f, force=True)
                    if str(getattr(ds, "StudyInstanceUID", "")) != uid or "PixelData" not in ds:
                        continue
                    arr = ds.pixel_array
                    if arr.ndim == 3 and arr.shape[-1] != 3:
                        arr = arr[arr.shape[0] // 2]
                    arr = arr * float(getattr(ds, "RescaleSlope", 1) or 1) + float(
                        getattr(ds, "RescaleIntercept", 0) or 0)
                    render_thumbnail(arr).save(path)
                    return path
                except Exception:  # noqa: BLE001 - DICOM이 아니거나 디코딩 불가
                    continue
    return None


def _safe(text):
    return re.sub(r"[^\w.\-가-힣]+", "_", str(text))[:120] or "study"


# ─── 기록 모으기 ───

def collection_paths(store):
    """컬렉션 id → '상위 / 하위' 경로 이름"""
    out = {}
    for c in store.collections:
        names, current, guard = [], c, 0
        while current is not None and guard < 50:
            names.append(current["name"])
            current = store.collection(current.get("parent"))
            guard += 1
        out[c["id"]] = " / ".join(reversed(names))
    return out


def build_records(store, uids, include, measure_provider=None, want_thumbnails=True):
    paths = collection_paths(store)
    records = []
    for uid in uids:
        s = store.get(uid)
        if s is None:
            continue
        rec = {"uid": uid, "patient_name": s.get("patient_name", ""), "patient_id": s.get("patient_id", ""),
               "study_date": _date(s.get("study_date")), "description": s.get("description", ""),
               "modalities": ", ".join(s.get("modalities", [])), "series_count": s.get("series", ""),
               "images": s.get("images", ""), "folder": s.get("folder", ""),
               "tags": list(s.get("tags", [])),
               "collections": [paths[c] for c in s.get("collections", []) if c in paths],
               "note_text": s.get("note_text", ""), "note_md": _html_to_markdown(s.get("note_html", ""))
               or s.get("note_text", ""),
               "series_list": list(s.get("series_list", [])), "thumbnail": None, "measurements": []}
        if include.get("thumbnail") and want_thumbnails:
            rec["thumbnail"] = find_thumbnail(store, uid, s)
        if include.get("measurements") and measure_provider is not None:
            rec["measurements"] = measure_provider(uid)
        records.append(rec)
    return records


def study_title(rec, include):
    parts = []
    if include.get("patient"):
        parts.append(f"{rec['patient_name'] or '-'} ({rec['patient_id'] or '-'})")
    if include.get("study"):
        parts.append(f"{rec['study_date']} {rec['description']}".strip())
    return " — ".join(p for p in parts if p) or rec["uid"]


def table_columns(include):
    cols = []
    if include.get("patient"):
        cols += [("patient_name", "환자명"), ("patient_id", "Patient ID")]
    if include.get("study"):
        cols += [("study_date", "검사일"), ("description", "Description"), ("modalities", "모달리티"),
                 ("series_count", "시리즈"), ("images", "영상")]
    if include.get("tags"):
        cols.append(("tags", "태그"))
    cols.append(("collections", "컬렉션"))
    if include.get("notes"):
        cols.append(("note_text", "메모"))
    cols.append(("folder", "폴더"))
    return cols


def cell(rec, key):
    value = rec.get(key, "")
    if key == "tags":
        return " ".join(f"#{t}" for t in value)
    if key == "collections":
        return "; ".join(value)
    return "" if value is None else str(value)


MEASURE_HEADER = ["시리즈", "슬라이스", "종류", "이름", "값", "Mean", "SD", "Min", "Max"]


# ─── 형식별 쓰기 ───

def export(fmt, records, include, path, store=None):
    if fmt == "hwp":
        raise UnsupportedFormat(HWP_MESSAGE)
    writer = {"pdf": write_pdf, "docx": write_docx, "xlsx": write_xlsx, "csv": write_csv,
              "md": write_markdown}.get(fmt)
    if writer is not None:
        return writer(records, include, path)
    if fmt == "json":
        return write_json(store, [r["uid"] for r in records], path)
    if fmt in ("png_sheet", "jpg_sheet"):
        return write_contact_sheet(records, include, path, "JPEG" if fmt.startswith("jpg") else "PNG")
    if fmt in ("png_each", "jpg_each"):
        return write_each_image(records, include, path, "jpg" if fmt.startswith("jpg") else "png")
    raise UnsupportedFormat(f"알 수 없는 형식: {fmt}")


def write_csv(records, include, path):
    cols = table_columns(include)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["StudyInstanceUID"] + [title for _k, title in cols])
        for r in records:
            writer.writerow([r["uid"]] + [cell(r, k) for k, _t in cols])
    return path


def write_json(store, uids, path):
    data = store.to_dict()
    data["studies"] = {u: store.studies[u] for u in uids if u in store.studies}
    data["exported"] = datetime.datetime.now().isoformat(timespec="seconds")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    return path


def write_markdown(records, include, path):
    lines = [f"# DabbaView Library — {datetime.date.today().isoformat()}", "",
             f"스터디 {len(records)}개", ""]
    for r in records:
        lines += [f"## {study_title(r, include)}", ""]
        if include.get("study"):
            lines.append(f"- 모달리티: {r['modalities']} · 시리즈 {r['series_count']} · 영상 {r['images']}")
        if include.get("tags") and r["tags"]:
            lines.append("- 태그: " + " ".join(f"#{t}" for t in r["tags"]))
        if r["collections"]:
            lines.append("- 컬렉션: " + "; ".join(r["collections"]))
        lines.append(f"- 폴더: `{r['folder']}`")
        if include.get("thumbnail") and r["thumbnail"]:
            lines += ["", f"![thumbnail]({r['thumbnail']})"]
        if include.get("series") and r["series_list"]:
            lines += ["", "| # | 시리즈 | 모달리티 | 영상 |", "|---|---|---|---|"]
            lines += [f"| {x.get('number', '')} | {x.get('description', '')} | {x.get('modality', '')} | "
                      f"{x.get('images', '')} |" for x in r["series_list"]]
        if include.get("measurements") and r["measurements"]:
            lines += ["", "| " + " | ".join(MEASURE_HEADER) + " |", "|" + "---|" * len(MEASURE_HEADER)]
            lines += ["| " + " | ".join(str(v) for v in m) + " |" for m in r["measurements"]]
        if include.get("notes") and r["note_md"]:
            lines += ["", r["note_md"]]
        lines.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return path


def write_xlsx(records, include, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Studies"
    cols = table_columns(include)
    header = [title for _k, title in cols] + ["StudyInstanceUID"]
    ws.append(header)
    for r in records:
        ws.append([cell(r, k) for k, _t in cols] + [r["uid"]])
    by_collection = wb.create_sheet("By Collection")
    by_collection.append(["컬렉션", "환자명", "Patient ID", "검사일", "Description", "태그"])
    for r in records:
        for c in r["collections"] or ["(컬렉션 없음)"]:
            by_collection.append([c, r["patient_name"], r["patient_id"], r["study_date"], r["description"],
                                  " ".join(f"#{t}" for t in r["tags"])])
    sheets = [ws, by_collection]
    if include.get("series"):
        series = wb.create_sheet("Series")
        series.append(["환자명", "검사일", "Study Description", "#", "Series Description", "모달리티", "영상"])
        for r in records:
            for x in r["series_list"]:
                series.append([r["patient_name"], r["study_date"], r["description"], x.get("number", ""),
                               x.get("description", ""), x.get("modality", ""), x.get("images", "")])
        sheets.append(series)
    if include.get("measurements"):
        meas = wb.create_sheet("Measurements")
        meas.append(["환자명", "검사일", "Study Description"] + MEASURE_HEADER)
        for r in records:
            for m in r["measurements"]:
                meas.append([r["patient_name"], r["study_date"], r["description"]] + list(m))
        sheets.append(meas)
    fill = PatternFill("solid", fgColor="DDE6F0")
    for sheet in sheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
            c.fill = fill
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = max(len(str(c.value or "")) for c in column)
            sheet.column_dimensions[column[0].column_letter].width = min(60, max(8, width * 1.2 + 2))
            for c in column[1:]:
                c.alignment = Alignment(vertical="top", wrap_text=width > 60)
    wb.save(path)
    return path


def write_docx(records, include, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Apple SD Gothic Neo" if sys.platform == "darwin" else "Malgun Gothic"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
    style.font.size = Pt(10)
    doc.add_heading(f"DabbaView Library — {datetime.date.today().isoformat()}", 0)
    doc.add_paragraph(f"스터디 {len(records)}개")
    cols = [c for c in table_columns(include) if c[0] not in ("note_text", "folder")]
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Light Grid Accent 1"
    for i, (_k, title) in enumerate(cols):
        table.rows[0].cells[i].text = title
    for r in records:
        row = table.add_row().cells
        for i, (k, _t) in enumerate(cols):
            row[i].text = cell(r, k)
    for r in records:
        doc.add_page_break() if len(records) > 1 and r is not records[0] else None
        doc.add_heading(study_title(r, include), level=1)
        info = []
        if include.get("study"):
            info.append(f"모달리티 {r['modalities']} · 시리즈 {r['series_count']} · 영상 {r['images']}")
        if include.get("tags") and r["tags"]:
            info.append("태그: " + " ".join(f"#{t}" for t in r["tags"]))
        if r["collections"]:
            info.append("컬렉션: " + "; ".join(r["collections"]))
        info.append("폴더: " + r["folder"])
        for line in info:
            doc.add_paragraph(line)
        if include.get("thumbnail") and r["thumbnail"]:
            doc.add_picture(r["thumbnail"], width=Cm(6))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
        if include.get("series") and r["series_list"]:
            doc.add_heading("시리즈", level=2)
            t = doc.add_table(rows=1, cols=4)
            t.style = "Light List Accent 1"
            for i, title in enumerate(["#", "Series Description", "모달리티", "영상"]):
                t.rows[0].cells[i].text = title
            for x in r["series_list"]:
                c = t.add_row().cells
                for i, key in enumerate(("number", "description", "modality", "images")):
                    c[i].text = str(x.get(key, ""))
        if include.get("measurements") and r["measurements"]:
            doc.add_heading("측정 / ROI", level=2)
            t = doc.add_table(rows=1, cols=len(MEASURE_HEADER))
            t.style = "Light List Accent 1"
            for i, title in enumerate(MEASURE_HEADER):
                t.rows[0].cells[i].text = title
            for m in r["measurements"]:
                c = t.add_row().cells
                for i, v in enumerate(m):
                    c[i].text = str(v)
        if include.get("notes") and r["note_md"]:
            doc.add_heading("메모", level=2)
            for line in r["note_md"].splitlines():
                stripped = line.strip()
                if not stripped:
                    continue
                if stripped.startswith("- "):
                    p = doc.add_paragraph(style="List Bullet")
                    stripped = stripped[2:]
                elif re.match(r"\d+\. ", stripped):
                    p = doc.add_paragraph(style="List Number")
                    stripped = stripped.split(". ", 1)[1]
                else:
                    p = doc.add_paragraph()
                for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", stripped):
                    if part.startswith("**") and part.endswith("**"):
                        p.add_run(part[2:-2]).bold = True
                    elif part.startswith("*") and part.endswith("*") and len(part) > 1:
                        p.add_run(part[1:-1]).italic = True
                    elif part:
                        p.add_run(part)
    doc.save(path)
    return path


def _pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    path = korean_font_path()
    if path:
        try:
            pdfmetrics.registerFont(TTFont("DVKorean", path))
            return "DVKorean"
        except Exception:  # noqa: BLE001
            pass
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    return "HYGothic-Medium"


def _pdf_markup(md_text):
    """메모 Markdown → reportlab 문단 마크업 (굵게·기울임·줄바꿈·목록)"""
    from xml.sax.saxutils import escape
    lines = []
    for line in md_text.splitlines():
        s = escape(line.strip())
        s = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", s)
        s = re.sub(r"\*([^*]+)\*", r"<i>\1</i>", s)
        if s.startswith("- "):
            s = "• " + s[2:]
        lines.append(s)
    return "<br/>".join(lines)


def write_pdf(records, include, path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    from xml.sax.saxutils import escape
    font = _pdf_font()
    base = ParagraphStyle("base", fontName=font, fontSize=9, leading=12)
    h1 = ParagraphStyle("h1", parent=base, fontSize=16, leading=20, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=base, fontSize=12, leading=16, spaceBefore=6, spaceAfter=4)
    small = ParagraphStyle("small", parent=base, fontSize=7.5, leading=9.5, textColor=colors.grey)

    def grid(data, widths=None):
        t = Table([[Paragraph(escape(str(c)), base) for c in row] for row in data], colWidths=widths,
                  repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE6F0")),
                               ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    story = [Paragraph(f"DabbaView Library — {datetime.date.today().isoformat()}", h1),
             Paragraph(f"스터디 {len(records)}개", base), Spacer(1, 4 * mm)]
    cols = [c for c in table_columns(include) if c[0] not in ("note_text", "folder")]
    story.append(grid([[t for _k, t in cols]] + [[cell(r, k) for k, _t in cols] for r in records]))
    for r in records:
        story += [PageBreak(), Paragraph(escape(study_title(r, include)), h1)]
        info = []
        if include.get("study"):
            info.append(f"모달리티 {r['modalities']} · 시리즈 {r['series_count']} · 영상 {r['images']}")
        if include.get("tags") and r["tags"]:
            info.append("태그: " + " ".join(f"#{t}" for t in r["tags"]))
        if r["collections"]:
            info.append("컬렉션: " + "; ".join(r["collections"]))
        for line in info:
            story.append(Paragraph(escape(line), base))
        story.append(Paragraph(escape("폴더: " + r["folder"]), small))
        if include.get("thumbnail") and r["thumbnail"]:
            story += [Spacer(1, 2 * mm), Image(r["thumbnail"], width=55 * mm, height=55 * mm)]
        if include.get("series") and r["series_list"]:
            story.append(Paragraph("시리즈", h2))
            story.append(grid([["#", "Series Description", "모달리티", "영상"]] +
                              [[x.get("number", ""), x.get("description", ""), x.get("modality", ""),
                                x.get("images", "")] for x in r["series_list"]],
                              [12 * mm, 100 * mm, 25 * mm, 20 * mm]))
        if include.get("measurements") and r["measurements"]:
            story.append(Paragraph("측정 / ROI", h2))
            story.append(grid([MEASURE_HEADER] + [list(m) for m in r["measurements"]]))
        if include.get("notes") and r["note_md"]:
            story += [Paragraph("메모", h2), Paragraph(_pdf_markup(r["note_md"]), base)]
    SimpleDocTemplate(path, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm,
                      bottomMargin=15 * mm, title="DabbaView Library").build(story)
    return path


def _label_lines(rec, include):
    lines = []
    if include.get("patient"):
        lines.append(f"{rec['patient_name'] or '-'} ({rec['patient_id'] or '-'})")
    if include.get("study"):
        lines.append(f"{rec['study_date']}  {rec['description']}".strip())
    if include.get("tags") and rec["tags"]:
        lines.append(" ".join(f"#{t}" for t in rec["tags"][:4]))
    return lines or [rec["uid"][-20:]]


def _tile(rec, include, size=256):
    from PIL import Image, ImageDraw
    thumb = Image.open(rec["thumbnail"]).convert("RGB") if rec.get("thumbnail") else None
    font = _pil_font(13)
    lines = _label_lines(rec, include)
    band = 8 + 18 * len(lines)
    tile = Image.new("RGB", (size, size + band), (24, 24, 24))
    if thumb is not None:
        tile.paste(thumb.resize((size, size)), (0, 0))
    else:
        ImageDraw.Draw(tile).text((size // 2 - 40, size // 2), "(썸네일 없음)", fill=(120, 120, 120), font=font)
    draw = ImageDraw.Draw(tile)
    for i, line in enumerate(lines):
        draw.text((6, size + 4 + 18 * i), line[:34], fill=(230, 230, 230), font=font)
    return tile


def write_contact_sheet(records, include, path, fmt="PNG"):
    from PIL import Image, ImageDraw
    if not records:
        raise UnsupportedFormat("내보낼 스터디가 없습니다.")
    tiles = [_tile(r, include) for r in records]
    cols = min(len(tiles), max(1, math.ceil(math.sqrt(len(tiles)))))
    rows = math.ceil(len(tiles) / cols)
    tw, th = tiles[0].size
    gap, top = 8, 40
    sheet = Image.new("RGB", (cols * (tw + gap) + gap, rows * (th + gap) + gap + top), (10, 10, 10))
    ImageDraw.Draw(sheet).text((gap, 10), f"DabbaView Library — {datetime.date.today().isoformat()} "
                               f"({len(records)} studies)", fill=(255, 255, 255), font=_pil_font(18))
    for i, tile in enumerate(tiles):
        r, c = divmod(i, cols)
        sheet.paste(tile, (gap + c * (tw + gap), top + gap + r * (th + gap)))
    sheet.save(path, fmt, **({"quality": 92} if fmt == "JPEG" else {}))
    return path


def write_each_image(records, include, folder, ext="png"):
    os.makedirs(folder, exist_ok=True)
    written = []
    for r in records:
        name = _safe(f"{r['study_date']}_{r['patient_name']}_{r['description']}") + f".{ext}"
        path = os.path.join(folder, name)
        tile = _tile(r, include)
        tile.save(path, "JPEG" if ext == "jpg" else "PNG", **({"quality": 92} if ext == "jpg" else {}))
        written.append(path)
    return folder


def preview_html(records, include, limit=15):
    """대화상자 미리보기 (표 + 메모 앞부분)"""
    from html import escape
    cols = [c for c in table_columns(include) if c[0] != "folder"]
    rows = "".join("<tr>" + "".join(f"<td>{escape(cell(r, k))[:120]}</td>" for k, _t in cols) + "</tr>"
                   for r in records[:limit])
    head = "".join(f"<th>{escape(t)}</th>" for _k, t in cols)
    more = f"<p>… 외 {len(records) - limit}개</p>" if len(records) > limit else ""
    return (f"<p><b>스터디 {len(records)}개</b></p>"
            f"<table border='1' cellspacing='0' cellpadding='3'><tr>{head}</tr>{rows}</table>{more}")
